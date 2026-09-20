import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl

paths = [
    Path(r"D:\桌面\android_doc\problem_record_filled_corrected.xlsx"),
    Path(r"D:\桌面\android_doc\problem_record.xlsx"),
    Path(r"D:\桌面\android_doc\problem_record_app_tracker.xlsx"),
]
out = Path(r"D:\桌面\Android_Automotive_Voice_Agent\scripts\_xlsx_dump.txt")
lines = []
for path in paths:
    lines.append("=" * 80)
    lines.append(str(path))
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        lines.append(f"--- sheet: {sheet} dims={ws.dimensions} ---")
        rows = list(ws.iter_rows(values_only=True))
        for i, row in enumerate(rows):
            vals = [("" if c is None else str(c).replace("\n", " | ")) for c in row]
            if not any(v.strip() for v in vals):
                continue
            lines.append(f"{i:03d}| " + " || ".join(vals))
out.write_text("\n".join(lines), encoding="utf-8")
print(f"wrote {out} ({len(lines)} lines)")
